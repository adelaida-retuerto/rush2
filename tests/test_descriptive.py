"""Tests des statistiques descriptives (src/pharma/analysis/descriptive.py).   [Ihman]

Lancer depuis la racine du projet :  python -m pytest tests/test_descriptive.py
On fabrique des ventes dont on connaît la saison, puis on vérifie que les fonctions la retrouvent,
et que chaque tableau a exactement les colonnes du modèle Excel.
"""
import sys
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import pandas as pd  # noqa: E402
import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from pharma import config  # noqa: E402
from pharma.analysis import descriptive  # noqa: E402

G = config.GROUP_CODES


@pytest.fixture
def c():
    """Ventes fabriquées de janvier 2017 à septembre 2019 (2 années complètes + 9 mois).

    Tous les groupes vendent 100 par mois, sauf en janvier où ils vendent 200 ;
    R06 vend en plus 300 en mai. 2019 vend 10 % de plus que 2018.
    """
    mois = pd.date_range("2017-01-31", "2019-09-30", freq="ME")
    valeurs = {}
    for g in G:
        serie = []
        for m in mois:
            v = 200.0 if m.month == 1 else 100.0
            if g == "R06" and m.month == 5:
                v = 300.0
            if m.year == 2019:
                v *= 1.1
            serie.append(v)
        valeurs[g] = serie
    monthly = pd.DataFrame(valeurs, index=mois)
    jours = pd.date_range("2019-09-01", "2019-09-30", freq="D")
    daily = pd.DataFrame({g: 2.0 for g in G}, index=jours)
    daily["N02BE"] = 10.0                                     # le paracétamol pèse plus que les autres
    return SimpleNamespace(daily=daily, monthly=monthly, last_complete_month=mois.max())


def colonnes_du_modele(tableau):
    """En-têtes d'un tableau du modèle Excel (templates/workbook_template.xlsx)."""
    wb = load_workbook(config.TEMPLATE_PATH)
    for ws in wb.worksheets:
        if tableau in ws.tables:
            debut, fin = ws.tables[tableau].ref.split(":")
            return [cellule.value for cellule in ws[debut:fin][0]]
    raise AssertionError(f"{tableau} absent du modèle")


@pytest.mark.parametrize("fonction, tableau", [
    ("group_volumes", "T_Groupes"), ("seasonality_index", "T_Saisonnalite"),
    ("peak_months", "T_Mois_Forts"), ("year_to_date", "T_Cumul"),
    ("yearly", "T_Annees"), ("monthly_table", "T_Mensuel"),
])
def test_colonnes_identiques_au_modele(c, fonction, tableau):
    resultat = getattr(descriptive, fonction)(c)
    assert list(resultat.columns) == colonnes_du_modele(tableau)


def test_group_volumes_parts(c):
    v = descriptive.group_volumes(c).set_index("Groupe")
    assert v["Part du volume"].sum() == pytest.approx(1, abs=0.001)
    assert v["Part du volume"].idxmax() == "N02BE"
    assert v.loc["N02BE", "Moyenne par jour"] == 10


def test_family_volumes(c):
    f = descriptive.family_volumes(c).set_index("Famille")
    assert set(f.index) == {"Anti-inflammatoires", "Antalgiques", "Psycholeptiques", "Respiratoire"}
    assert f.loc["Antalgiques", "Groupes"] == "N02BA, N02BE"
    assert f.loc["Antalgiques", "Moyenne par jour"] == 12          # aspirine 2 + paracétamol 10
    assert f["Part du volume"].sum() == pytest.approx(1, abs=0.001)


def test_seasonality_retrouve_janvier_et_mai(c):
    s = descriptive.seasonality_index(c).set_index("Mois")
    assert len(s) == 12
    # Sur les années complètes (2017-2018), janvier vaut 2 fois les autres mois
    assert s.loc["Janvier", "M01AB"] == pytest.approx(2 * s.loc["Février", "M01AB"], rel=0.01)
    assert s["N02BA"].mean() == pytest.approx(1, abs=0.01)  # indice moyen = 1
    assert s["R06"].idxmax() == "Mai"


def test_peak_months(c):
    p = descriptive.peak_months(c).set_index("Groupe")
    assert p.loc["R06", "Mois le plus fort"] == "Mai"
    assert p.loc["M01AB", "Mois le plus fort"] == "Janvier"
    assert p.loc["M01AB", "Amplitude max sur min"] == pytest.approx(2, abs=0.01)


def test_year_to_date_meme_periode(c):
    cumul = descriptive.year_to_date(c).set_index("Groupe")
    # Janvier à septembre : 2019 vend 10 % de plus que 2018
    assert cumul.loc["M01AB", "Évolution"] == pytest.approx(0.10, abs=0.001)
    assert cumul.attrs["periode"] == "janvier à septembre 2019 face à 2018"


def test_yearly_seulement_annees_completes(c):
    annees = descriptive.yearly(c)
    assert list(annees["Année"]) == [2017, 2018]            # 2019 n'a que 9 mois
    assert annees.loc[0, "Total"] == pytest.approx(annees[G].iloc[0].sum())


def test_monthly_table_mois_en_date(c):
    m = descriptive.monthly_table(c)
    assert len(m) == len(c.monthly)
    assert m["Mois"].iloc[-1] == pd.Timestamp("2019-09-01")  # une vraie date, pour le graphique
