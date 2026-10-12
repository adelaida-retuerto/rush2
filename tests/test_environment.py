"""Tests de l'analyse grippe (src/pharma/analysis/environment.py).   [Ihman]

Lancer depuis la racine du projet :  python -m pytest tests/test_environment.py
Aucun accès à internet : on travaille sur un petit fichier Sentinelles fabriqué, ou hors ligne.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from pharma import config  # noqa: E402
from pharma.analysis import environment  # noqa: E402

G = config.GROUP_CODES


def colonnes_du_modele(tableau):
    """En-têtes d'un tableau du modèle Excel (templates/workbook_template.xlsx)."""
    wb = load_workbook(config.TEMPLATE_PATH)
    ws = wb["Environnement"]
    debut, fin = ws.tables[tableau].ref.split(":")
    return [cellule.value for cellule in ws[debut:fin][0]]


def test_load_flu_lit_et_date_les_semaines(tmp_path, monkeypatch):
    """Une copie locale au format Sentinelles : semaines ISO datées du dimanche, « - » devient vide."""
    fichier = tmp_path / "sentiweb.csv"
    fichier.write_text(
        "# commentaire du réseau Sentinelles\n"
        "week,indicator,inc,inc_low,inc_up,inc100,inc100_low,inc100_up,geo_insee,geo_name\n"
        "201404,3,1000,900,1100,150,140,160,FR,France\n"
        "201403,3,800,700,900,120,110,130,FR,France\n"
        "198919,3,-,,,-,,,FR,France\n",
        encoding="utf-8")
    monkeypatch.setattr(config, "SENTIWEB_CACHE", fichier)

    flu, source = environment.load_flu(offline=True)
    assert "copie locale" in source
    assert flu.index.is_monotonic_increasing                       # du plus ancien au plus récent
    assert flu.loc[pd.Timestamp("2014-01-19"), "incidence"] == 120   # semaine 3 de 2014 = dimanche 19/01
    assert (flu.index.dayofweek == 6).all()
    assert flu["incidence"].isna().sum() == 1                       # la semaine notée « - »


def test_load_flu_sans_copie_locale(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SENTIWEB_CACHE", tmp_path / "absent.csv")
    with pytest.raises(FileNotFoundError):
        environment.load_flu(offline=True)


@pytest.fixture
def donnees_fabriquees():
    """4 ans de semaines : N02BE = effet du mois + 0,5 vente par cas de grippe ; R06 ne dépend que du mois."""
    semaines = pd.date_range("2015-01-04", "2018-12-30", freq="W-SUN")
    rng = np.random.default_rng(0)
    # Grippe : pic chaque hiver, d'une hauteur différente selon l'année
    hauteur = {2015: 300, 2016: 100, 2017: 500, 2018: 200}
    incidence = [hauteur[d.year] if d.month in (1, 2) else 10 for d in semaines]
    flu = pd.DataFrame({"incidence": incidence}, index=semaines)
    base = {g: 50 + 10 * semaines.month.isin([12, 1, 2]).astype(float) + rng.normal(0, 1, len(semaines)) for g in G}
    weekly = pd.DataFrame(base, index=semaines)
    weekly["N02BE"] += 0.5 * flu["incidence"].values
    return weekly, flu


def test_flu_effect_colonnes_identiques_au_modele(donnees_fabriquees):
    weekly, flu = donnees_fabriquees
    effet, serie = environment.flu_effect(weekly, flu)
    assert list(effet.columns) == colonnes_du_modele("T_Grippe_Effet")
    assert list(serie.columns) == colonnes_du_modele("T_Grippe_Serie")
    assert len(serie) <= config.CHART_WINDOW_WEEKS
    assert effet.attrs["n_weeks"] > 0 and "threshold" in effet.attrs


def test_flu_effect_retrouve_un_vrai_effet(donnees_fabriquees):
    """N02BE a été fabriqué avec +0,5 vente par cas : l'effet par +100 cas doit valoir environ 50."""
    weekly, flu = donnees_fabriquees
    effet = environment.flu_effect(weekly, flu)[0].set_index("Groupe")
    assert effet.loc["N02BE", "Effet par +100 cas/100k (même mois)"] == pytest.approx(50, rel=0.1)
    assert effet.loc["N02BE", "Lien significatif (p<0,01)"] == "Oui"
    # R06 n'a reçu aucun effet de la grippe : son effet doit rester proche de 0 (bien loin de 50).
    # On teste la taille de l'effet, pas sa p-valeur : avec 4 hivers fabriqués, le hasard seul
    # peut donner une p-valeur basse (c'est la fragilité discutée dans l'analyse).
    assert abs(effet.loc["R06", "Effet par +100 cas/100k (même mois)"]) < 5


def test_semaine_precedente_decalee_dune_semaine(donnees_fabriquees):
    weekly, flu = donnees_fabriquees
    donnees = environment.aligner(weekly, flu)
    # L'incidence S-1 d'une semaine = l'incidence de la semaine d'avant
    assert donnees["incidence_s1"].iloc[5] == donnees["incidence"].iloc[4]


def test_link_ladder(donnees_fabriquees):
    weekly, flu = donnees_fabriquees
    grille = environment.link_ladder(environment.flu_effect(weekly, flu)[0])
    assert list(grille.columns) == colonnes_du_modele("T_Grille_Liens")
    assert len(grille) == 5
    assert grille.iloc[-1]["Niveau"] == "5. Causalité"
    assert "0." not in grille.iloc[0]["Exemple dans ce projet"]      # nombres à la française


def test_vraies_donnees_hors_ligne():
    """Avec la copie Sentinelles du dépôt et les vraies ventes : 300 semaines, aucun lien significatif."""
    from pharma.cleaning import hourly, loading
    from pharma.cleaning.weekly import clean_weekly
    journal = []
    h = hourly.clean_hourly(loading.load_export("hourly"), loading.load_derived_columns("hourly"), journal)
    weekly, _ = clean_weekly(loading.load_export("weekly"), None, h.daily_all, h.complete_day, journal)
    flu, _ = environment.load_flu(offline=True)
    effet, _ = environment.flu_effect(weekly, flu)
    assert effet.attrs["n_weeks"] == 300
    assert (effet["Lien significatif (p<0,01)"] == "Non").all()
    assert effet.set_index("Groupe").loc["N02BE", "Hiver : ventes sem. forte grippe / autres"] > 1.1
