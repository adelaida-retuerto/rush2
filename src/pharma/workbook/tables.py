"""Outils communs du classeur : déclarer les zones du modèle, puis les remplir.   [Adelaida]

Les 4 règles de la mise à jour mensuelle sont appliquées ici :

1. On part toujours du modèle (`templates/workbook_template.xlsx`), jamais du fichier du mois précédent.
2. Le code n'écrit QUE dans deux types de zones déclarées dans le modèle :
     - des tableaux Excel nommés (ex. T_Jour), qui reçoivent les données ;
     - des cellules nommées (ex. periode_debut), qui reçoivent un texte ou une valeur.
   Titres, formules de l'outil, graphiques et mises en forme ne sont jamais touchés.
3. Les formules pointent sur les tableaux (T_Jour[Date]) : quand un tableau grandit, elles suivent.
4. Après la mise à jour, `verify.py` contrôle que rien d'autre n'a bougé.
"""
from __future__ import annotations

import datetime as dt
from copy import copy

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import styles as S


class TemplateError(RuntimeError):
    """Le modèle ne correspond pas au contrat (zone absente, colonnes différentes, chevauchement)."""


# ============================================================================ construction du modèle

def title(ws, text: str, cell: str = "B2") -> None:
    ws[cell] = text
    ws[cell].font = S.TITLE


def note(ws, cell: str, text: str) -> None:
    ws[cell] = text
    ws[cell].font = S.NOTE


def h2(ws, cell: str, text: str) -> None:
    ws[cell] = text
    ws[cell].font = S.H2


def widths(ws, values: dict[str, float]) -> None:
    for col, w in values.items():
        ws.column_dimensions[col].width = w


def add_table(ws, name: str, row: int, col: int, headers: list[str], *, capacity: int = 1,
              heading: str | None = None, comment: str | None = None,
              formats: dict[str, str] | None = None, wrap: bool = False) -> tuple[int, int]:
    """Déclare un tableau Excel vide dans le modèle.

    `capacity` = nombre de lignes réservées sous le tableau avant la zone suivante (taille fixe connue).
    Renvoie (ligne d'en-tête, première ligne libre après la zone réservée).
    """
    formats = formats or {}
    if heading:
        ws.cell(row, col, heading).font = S.H2
        row += 1
    if comment:
        ws.cell(row, col, comment).font = S.NOTE
        row += 1
    for j, h in enumerate(headers):
        c = ws.cell(row, col + j, h)
        c.font, c.fill, c.border, c.alignment = S.HEAD, S.HEAD_FILL, S.BOX, S.HEAD_ALIGN
        d = ws.cell(row + 1, col + j)  # ligne modèle : son style sera recopié sur chaque ligne de données
        d.font, d.border = S.BODY, S.BOX
        d.number_format = formats.get(h, S.guess_format(h))
        if wrap:
            d.alignment = S.WRAP
    ws.row_dimensions[row].height = 30
    ref = f"{get_column_letter(col)}{row}:{get_column_letter(col + len(headers) - 1)}{row + 1}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=S.TABLE_STYLE, showRowStripes=False)
    ws.add_table(table)
    return row, row + 1 + capacity + 2


def name_cell(wb, name: str, ws, coord: str, value=None, font=None) -> None:
    """Déclare une cellule nommée que le code a le droit de remplir."""
    if value is not None:
        ws[coord] = value
    if font is not None:
        ws[coord].font = font
    col, row = "".join(ch for ch in coord if ch.isalpha()), "".join(ch for ch in coord if ch.isdigit())
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{ws.title}'!${col}${row}")


# ============================================================================ lecture du modèle

def tables(wb) -> dict[str, tuple]:
    """{nom du tableau: (feuille, objet Table)} pour tout le classeur."""
    out = {}
    for ws in wb.worksheets:
        for t in ws.tables.values():
            out[t.displayName] = (ws, t)
    return out


def named_cells(wb) -> dict[str, tuple[str, str]]:
    out = {}
    for name, dn in wb.defined_names.items():
        for sheet, coord in dn.destinations:
            out[name] = (sheet, coord.replace("$", ""))
    return out


def check_contract(wb, expected_tables: list[str], expected_names: list[str]) -> None:
    """Arrête tout si le modèle a perdu une zone attendue (ex. tableau supprimé à la main)."""
    t, n = tables(wb), named_cells(wb)
    missing = [x for x in expected_tables if x not in t] + [x for x in expected_names if x not in n]
    if missing:
        raise TemplateError(f"Zones absentes du modèle : {missing}. Recréez-le avec scripts/build_template.py")


# ============================================================================ remplissage mensuel

def _clean_value(v):
    if v is None:
        return None
    if isinstance(v, float) and np.isnan(v):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, np.generic):
        return v.item()
    if isinstance(v, (dt.date, dt.datetime, str, int, float, bool)):
        return v
    return str(v)


def _overlaps(a: str, b: str) -> bool:
    a1, a2, a3, a4 = range_boundaries(a)
    b1, b2, b3, b4 = range_boundaries(b)
    return not (a3 < b1 or b3 < a1 or a4 < b2 or b4 < a2)


def fill_table(wb, name: str, df: pd.DataFrame) -> int:
    """Remplit un tableau du modèle avec `df` et l'agrandit ; renvoie le nombre de lignes écrites."""
    found = tables(wb).get(name)
    if found is None:
        raise TemplateError(f"Tableau {name} absent du modèle")
    ws, table = found
    c1, r1, c2, r2 = range_boundaries(table.ref)
    headers = [ws.cell(r1, c).value for c in range(c1, c2 + 1)]
    if [str(h) for h in headers] != [str(h) for h in df.columns]:
        raise TemplateError(f"{name} : colonnes du modèle {headers} différentes des données {list(df.columns)}")

    row_styles = [copy(ws.cell(r1 + 1, c)._style) for c in range(c1, c2 + 1)]
    heights = ws.row_dimensions[r1 + 1].height
    n = max(len(df), 1)                      # un tableau Excel garde toujours au moins une ligne
    new_ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r1 + n}"
    for other_ws, other in tables(wb).values():
        if other_ws is ws and other.displayName != name and _overlaps(new_ref, other.ref):
            raise TemplateError(f"{name} ({new_ref}) déborderait sur {other.displayName} ({other.ref}) : "
                                "agrandir l'espace réservé dans le modèle")

    values = df.to_numpy(dtype=object)
    for i in range(n):
        for j in range(c2 - c1 + 1):
            cell = ws.cell(r1 + 1 + i, c1 + j)
            cell.value = _clean_value(values[i, j]) if i < len(df) else None
            cell._style = copy(row_styles[j])
        if heights:
            ws.row_dimensions[r1 + 1 + i].height = heights
    table.ref = new_ref
    if table.autoFilter is not None:
        table.autoFilter.ref = new_ref
    return len(df)


def set_named(wb, name: str, value) -> None:
    """Écrit une valeur dans une cellule nommée du modèle (le style de la cellule est conservé)."""
    cells = named_cells(wb)
    if name not in cells:
        raise TemplateError(f"Cellule nommée {name} absente du modèle")
    sheet, coord = cells[name]
    wb[sheet][coord].value = _clean_value(value)
