"""Feuilles Stats_Groupes, Stats_Mensuel et Environnement.   [Ihman]

fill(wb, c, t) : remplit chaque mois les tableaux et cellules nommées du modèle (lancé par main.py).
build(wb) : dessine les feuilles dans le modèle templates/workbook_template.xlsx (titres, graphiques,
emplacement des tableaux) ; lancé une fois, par template.build_template.
"""
from __future__ import annotations

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

from .. import config
from . import styles as S
from . import tables as X

G = config.GROUP_CODES
SHEETS = ["Stats_Groupes", "Stats_Mensuel", "Environnement"]

# Les zones que ce module a le droit de remplir (contrôlées par template.contract et verify)
TABLES = ["T_Groupes", "T_Saisonnalite", "T_Mois_Forts", "T_Cumul", "T_Annees", "T_Mensuel",
          "T_Grille_Liens", "T_Grippe_Effet", "T_Grippe_Serie"]
NAMES = ["txt_cumul_titre", "txt_grippe_source", "txt_grippe_note"]


# Titres des colonnes de chaque tableau (identiques aux tables calculées par descriptive.py et environment.py)
H_GROUPES = ["Groupe", "Libellé", "Total période", "Part du volume", "Moyenne par jour", "Médiane par jour"]
H_SAISON = ["Mois", *G]
H_MOIS_FORTS = ["Groupe", "Mois le plus fort", "Indice max", "Mois le plus faible", "Indice min",
                "Amplitude max sur min"]
H_CUMUL = ["Groupe", "Cumul année en cours", "Cumul année précédente", "Évolution"]
H_ANNEES = ["Année", *G, "Total"]
H_MENSUEL = ["Mois", *G, "Total"]
H_GRILLE = ["Niveau", "Définition", "Exemple dans ce projet", "Ce qu'on peut en conclure"]
H_EFFET = ["Groupe", "Corrélation de Spearman (brute)", "Effet par +100 cas/100k (même mois)", "p-valeur",
           "Effet en % de la moyenne hebdo", "R² saison seule", "R² saison + grippe", "Gain de R²",
           "Effet avec grippe S-1", "p-valeur S-1", "Hiver : ventes sem. forte grippe / autres",
           "Lien significatif (p<0,01)", "Nature du lien", "Causalité"]
H_SERIE = ["Semaine", *G, "Incidence grippe"]


def _sheet(wb, name, color):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = color
    ws.column_dimensions["A"].width = 2
    return ws


def _window(ws, table: str, columns: list[str], size: int, top: int, left: int, formats: list[str]) -> None:
    """Fenêtre glissante : les `size` dernières lignes de `table`, en formules INDEX (taille fixe).

    Les tableaux mensuel et grippe grandissent chaque mois ; leurs graphiques lisent cette fenêtre,
    qui garde toujours la même taille : le graphique n'a jamais besoin d'être modifié.
    """
    for j, col in enumerate(columns):
        head = ws.cell(top, left + j, col)
        head.font, head.fill = S.HEAD, S.HEAD_FILL
        for k in range(1, size + 1):
            pos = f"ROWS({table}[{columns[0]}])-{size}+{k}"
            c = ws.cell(top + k, left + j, f'=IF({pos}<1,"",INDEX({table}[{col}],{pos}))')
            c.number_format, c.font = formats[j], S.NOTE


def build(wb) -> None:
    """Dessine les feuilles Stats_Groupes, Stats_Mensuel et Environnement dans le modèle."""
    # ------------------------------------------------------------------ Stats_Groupes
    ws = _sheet(wb, "Stats_Groupes", S.SAUGE)
    X.title(ws, "Que vend la pharmacie ? Poids des groupes et saisons")
    X.note(ws, "B3", f"Jours complets uniquement ; {config.UNIT_LABEL}.")
    X.widths(ws, {"B": 12, "C": 52, **{get_column_letter(c): 12 for c in range(4, 13)}})
    head, nxt = X.add_table(ws, "T_Groupes", 5, 2, H_GROUPES, capacity=8, heading="Poids de chaque groupe")
    ch = BarChart()
    ch.title, ch.height, ch.width, ch.legend = "Part du volume total", 7, 14, None
    ch.add_data(Reference(ws, min_col=5, min_row=head, max_row=head + 8), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=head + 1, max_row=head + 8))
    ch.series[0].graphicalProperties.solidFill = S.SAUGE
    ch.y_axis.number_format = "0%"
    ws.add_chart(ch, f"I{head}")

    head, nxt = X.add_table(ws, "T_Saisonnalite", nxt + 4, 2, H_SAISON, capacity=12,
                            heading="Saisonnalité : indice mensuel par groupe",
                            comment="Moyenne journalière du mois / moyenne journalière globale (1,50 = +50 %).",
                            formats={g: S.FMT_IDX for g in G})
    # Carte de chaleur : bleu = mois creux, blanc = mois habituel, rouge = mois fort
    ws.conditional_formatting.add(f"C{head + 1}:J{head + 12}", ColorScaleRule(
        start_type="num", start_value=0.5, start_color=S.COULEURS_ATC["R06"], mid_type="num", mid_value=1,
        mid_color="FFFFFF", end_type="num", end_value=1.7, end_color=S.COULEURS_ATC["N02BA"]))
    head, nxt = X.add_table(ws, "T_Mois_Forts", nxt, 2, H_MOIS_FORTS, capacity=8, heading="Mois forts et faibles")
    X.name_cell(wb, "txt_cumul_titre", ws, f"B{nxt}", "", S.H2)
    head, nxt = X.add_table(ws, "T_Cumul", nxt + 1, 2, H_CUMUL, capacity=9)
    X.add_table(ws, "T_Annees", nxt, 2, H_ANNEES, heading="Volumes par année civile complète",
                formats={"Année": "0"})

    # ------------------------------------------------------------------ Stats_Mensuel
    ws = _sheet(wb, "Stats_Mensuel", S.SAUGE)
    X.title(ws, "Série mensuelle recalculée depuis l'export horaire")
    X.note(ws, "B3", "Mois complets uniquement. Le graphique montre les 60 derniers mois.")
    X.widths(ws, {"B": 11, **{get_column_letter(c): 10 for c in range(3, 12)}, "L": 3})
    X.add_table(ws, "T_Mensuel", 5, 2, H_MENSUEL)
    w = config.CHART_WINDOW_MONTHS
    _window(ws, "T_Mensuel", ["Mois", *G], w, 5, 13, [S.FMT_MONTH] + [S.FMT_QTY] * len(G))
    lc = LineChart()
    lc.title, lc.height, lc.width = f"Ventes mensuelles par groupe ({w} derniers mois)", 10, 26
    lc.add_data(Reference(ws, min_col=14, max_col=13 + len(G), min_row=5, max_row=5 + w), titles_from_data=True)
    lc.set_categories(Reference(ws, min_col=13, min_row=6, max_row=5 + w))
    lc.x_axis.number_format = "mm/yy"
    for serie, couleur in zip(lc.series, S.PALETTE):  # chaque groupe garde sa couleur de la charte
        serie.graphicalProperties.line.solidFill = couleur
    ws.add_chart(lc, "W5")

    # ------------------------------------------------------------------ Environnement
    ws = _sheet(wb, "Environnement", S.COULEURS_ATC["R03"])
    X.title(ws, "Pharmacie ou environnement ? L'effet de la grippe")
    X.name_cell(wb, "txt_grippe_source", ws, "B3", "", S.NOTE)
    X.widths(ws, {"B": 24, "C": 30, "D": 34, "E": 34, **{get_column_letter(c): 15 for c in range(6, 17)}})
    head, nxt = X.add_table(ws, "T_Grille_Liens", 5, 2, H_GRILLE, capacity=4, wrap=True,
                            heading="Grille de lecture : association, levier actionnable, causalité",
                            comment="Exemples calculés sur le paracétamol (N02BE). Une association ne prouve pas une cause.")
    X.name_cell(wb, "txt_grippe_note", ws, f"B{nxt}", "", S.NOTE)
    head, nxt = X.add_table(ws, "T_Grippe_Effet", nxt + 1, 2, H_EFFET, capacity=8,
                            heading="Effet de l'incidence grippale sur les ventes hebdomadaires",
                            formats={"Lien significatif (p<0,01)": S.FMT_TEXT, "Nature du lien": S.FMT_TEXT})
    top = nxt
    X.add_table(ws, "T_Grippe_Serie", top, 2, H_SERIE, heading="Série hebdomadaire : ventes et grippe",
                formats={"Semaine": S.FMT_DATE, "Incidence grippe": S.FMT_INT})
    w = config.CHART_WINDOW_WEEKS
    _window(ws, "T_Grippe_Serie", ["Semaine", "N02BE", "Incidence grippe"], w, top + 1, 14,
            [S.FMT_DATE, S.FMT_QTY, S.FMT_INT])
    lc = LineChart()
    lc.title, lc.height, lc.width = "Paracétamol (N02BE) et incidence grippale, 3 dernières années", 9, 24
    lc.add_data(Reference(ws, min_col=15, min_row=top + 1, max_row=top + 1 + w), titles_from_data=True)
    lc.set_categories(Reference(ws, min_col=14, min_row=top + 2, max_row=top + 1 + w))
    lc.y_axis.title, lc.x_axis.number_format = "Quantité / semaine", "mm/yy"
    lc.series[0].graphicalProperties.line.solidFill = S.COULEURS_ATC["N02BE"]
    lc2 = LineChart()  # deuxième axe vertical, à droite, pour l'incidence
    lc2.add_data(Reference(ws, min_col=16, min_row=top + 1, max_row=top + 1 + w), titles_from_data=True)
    lc2.y_axis.axId, lc2.y_axis.title, lc2.y_axis.crosses = 200, "Cas / 100 000 hab.", "max"
    lc2.series[0].graphicalProperties.line.solidFill = S.GRIS
    lc += lc2
    ws.add_chart(lc, f"R{top + 1}")


def fill(wb, c, t) -> dict[str, int]:
    """Écrit les tableaux de statistiques et de grippe ; renvoie le nombre de lignes de chaque tableau.

    t : les tables calculées par pipeline.analyse_sales (descriptive.py et environment.py)
    """
    lignes = {
        # Feuille Stats_Groupes
        "T_Groupes": X.fill_table(wb, "T_Groupes", t["group_volumes"]),
        "T_Saisonnalite": X.fill_table(wb, "T_Saisonnalite", t["seasonality"]),
        "T_Mois_Forts": X.fill_table(wb, "T_Mois_Forts", t["peaks"]),
        "T_Cumul": X.fill_table(wb, "T_Cumul", t["year_to_date"]),
        "T_Annees": X.fill_table(wb, "T_Annees", t["yearly"]),
        # Feuille Stats_Mensuel
        "T_Mensuel": X.fill_table(wb, "T_Mensuel", t["monthly"]),
        # Feuille Environnement
        "T_Grille_Liens": X.fill_table(wb, "T_Grille_Liens", t["ladder"]),
        "T_Grippe_Effet": X.fill_table(wb, "T_Grippe_Effet", t["flu_effect"]),
        "T_Grippe_Serie": X.fill_table(wb, "T_Grippe_Serie", t["flu_series"]),
    }

    # Cellules de texte : titre du cumul (la période change chaque mois), source et note de la grippe
    X.set_named(wb, "txt_cumul_titre", f"Cumul de l'année : {t['year_to_date'].attrs['periode']}")
    X.set_named(wb, "txt_grippe_source", t["flu_source"])
    X.set_named(wb, "txt_grippe_note", t["flu_note"])
    return lignes
