"""Feuilles Synthèse et Prévision.   [Synthèse : Adelaida · Prévision : Jeanne]

La Synthèse est la page d'accueil du client : une lecture stratégique en une page.
  0. Message clé          : la situation en une phrase.
  1. Chiffres clés        : 6 cartes, un grand chiffre et une phrase courte.
  2. Ce qu'il faut retenir : 6 cartes d'insights (thème, chiffre marquant, message).
  3. Prévision du mois suivant : verdict par groupe (tableau, il grandit avec les groupes).
  4. Recommandations      : 5 leviers d'action, chacun précis et chiffré.
  5. Limites              : ce que les données ne permettent pas de conclure.
Les cartes sont des cellules nommées (taille fixe) ; le code n'écrit que dans ces cellules et
dans les tableaux déclarés, le reste du modèle (mise en forme) n'est jamais touché.

build(wb) : dessine les feuilles dans le modèle (une fois, par template.build_template).
fill(wb, c, t) : remplit les cartes et les tableaux chaque mois (par update.update_workbook).
"""
from __future__ import annotations

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import config
from . import styles as S
from . import tables as X

SHEETS = ["Synthèse", "Prévision"]

# --- Tailles fixes des blocs de la Synthèse (le code de summary.py produit exactement ces nombres)
N_KPI, N_INSIGHTS, N_RECOS, N_LIMITES = 6, 6, 5, 9

TABLES = ["T_Verdict", "T_Prevision", "T_Valeur", "T_Regularite", "T_Backtest", "T_Detail_Backtest"]
NAMES = (["txt_sous_titre", "txt_prevision_note", "message_cle"]
         + [f"kpi_{i}_{champ}" for i in range(1, N_KPI + 1) for champ in ("valeur", "libelle")]
         + [f"retenir_{i}_{champ}" for i in range(1, N_INSIGHTS + 1) for champ in ("theme", "chiffre", "texte")]
         + [f"reco_{i}_{champ}" for i in range(1, N_RECOS + 1) for champ in ("axe", "texte")]
         + [f"limite_{i}_{champ}" for i in range(1, N_LIMITES + 1) for champ in ("limite", "consequence", "levee")])

# Colonnes des tableaux de prévision : celles que remplit Jeanne (verdict_table et fill_* ci-dessous)
H_VERDICT = ["Groupe", "Meilleure référence naïve", "Erreur du modèle en % du niveau",
             "Gain du modèle vs meilleure référence", "Verdict : prévisible ?", "Prévision",
             "Fourchette basse (80 %)", "Fourchette haute (80 %)"]
H_PREVISION = ["Groupe", "Mois prévu", "Méthode retenue", "Prévision", "Fourchette basse (80 %)",
               "Fourchette haute (80 %)", "Verdict"]
H_VALEUR = ["Groupe", "Verdict : prévisible ?", "Méthode recommandée", "Erreur évitée par le modèle (unités/mois)",
            "Équivalent en jours de vente"]
H_REGULARITE = ["Groupe", "Libellé", "Moyenne par jour", "Écart-type par jour", "Coef. de variation",
                "P95 par jour", "Max par jour", "Part de jours sans vente"]
H_BACKTEST = ["Groupe", "Niveau moyen mensuel (période test)", "MAE Naïf (mois précédent)",
              "MAE Naïf saisonnier (même mois N-1)", "MAE Moyenne 3 derniers mois", "MAE Lissage exponentiel saisonnier",
              "MASE Naïf (mois précédent)", "MASE Naïf saisonnier (même mois N-1)", "MASE Moyenne 3 derniers mois",
              "MASE Lissage exponentiel saisonnier", "Meilleure référence naïve", "Erreur du modèle en % du niveau",
              "Gain du modèle vs meilleure référence", "Mois où le modèle fait mieux", "p-valeur (Wilcoxon unilatéral)",
              "Verdict : prévisible ?", "Erreur évitée par le modèle (unités/mois)", "Équivalent en jours de vente",
              "Méthode recommandée"]
H_DETAIL = ["Groupe", "Mois testé", "Méthode", "Prévision", "Réel", "Erreur absolue"]
TXT = {"Mois prévu": S.FMT_TEXT, "Méthode retenue": S.FMT_TEXT, "Verdict": S.FMT_TEXT,
       "Méthode recommandée": S.FMT_TEXT, "Méthode": S.FMT_TEXT, "Mois testé": S.FMT_TEXT,
       "Meilleure référence naïve": S.FMT_TEXT, "Verdict : prévisible ?": S.FMT_TEXT}

# --- Mise en forme des cartes (charte : sapin, sauge, brume ; ambre réservé aux alertes) -------
FOND_CARTE = PatternFill(fill_type="solid", start_color=S.BRUME, end_color=S.BRUME)
# Tailles de la Synthèse : titre 28, titres secondaires 18, texte 12 (lisible à l'écran comme à l'impression)
TITRE_SYNTHESE = Font(name=S.POLICE_TITRE, size=28, bold=True, color=S.SAPIN)
TITRE_SECONDAIRE = Font(name=S.POLICE_TITRE, size=18, bold=True, color=S.SAPIN)
MESSAGE = Font(name=S.POLICE_TITRE, size=18, italic=True, color=S.SAUGE)
GRAND_CHIFFRE = Font(name=S.POLICE_TITRE, size=28, bold=True, color=S.SAPIN)
CHIFFRE_ALERTE = Font(name=S.POLICE_TITRE, size=28, bold=True, color=S.AMBRE)
CHIFFRE_INSIGHT = Font(name=S.POLICE_TITRE, size=18, bold=True, color=S.SAPIN)
LIBELLE = Font(name=S.POLICE_TEXTE, size=12, color=S.ENCRE)
THEME = Font(name=S.POLICE_TEXTE, size=12, bold=True, color="FFFFFF")
TEXTE = Font(name=S.POLICE_TEXTE, size=12, color=S.ENCRE)
TEXTE_GRAS = Font(name=S.POLICE_TEXTE, size=12, bold=True, color=S.SAPIN)
NOTE_12 = Font(name=S.POLICE_TEXTE, size=12, italic=True, color=S.GRIS)
ENTETE_12 = Font(name=S.POLICE_TEXTE, size=12, bold=True, color="FFFFFF")
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
HAUT_GAUCHE = Alignment(horizontal="left", vertical="top", wrap_text=True)
GAUCHE = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
TRAIT = Side(style="thin", color="D9DEDC")
LIGNE_BAS = Border(bottom=TRAIT)

# Couleur du bandeau de chaque carte « Ce qu'il faut retenir » (une teinte de la charte par thème)
COULEURS_INSIGHTS = [S.SAPIN, S.COULEURS_ATC["N02BE"], S.COULEURS_ATC["R03"], S.SAUGE,
                     S.COULEURS_ATC["M01AE"], S.COULEURS_ATC["N05C"]]  # teintes foncées : texte blanc lisible
# La carte de chiffre clé n° 6 (partage des données) est une alerte : chiffre en ambre
KPI_ALERTE = {6}


def _sheet(wb, name, color):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = color
    ws.column_dimensions["A"].width = 2
    return ws


def _bloc(wb, ws, nom, ligne, col1, col2, font, align, fill=None, hauteur=None, border=None):
    """Fusionne une plage d'une ligne, la met en forme et la déclare comme cellule nommée."""
    debut, fin = f"{get_column_letter(col1)}{ligne}", f"{get_column_letter(col2)}{ligne}"
    if col2 > col1:
        ws.merge_cells(f"{debut}:{fin}")
    for col in range(col1, col2 + 1):  # le fond et la bordure doivent couvrir toute la plage fusionnée
        cell = ws.cell(ligne, col)
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border
    ws[debut].font, ws[debut].alignment = font, align
    if hauteur:
        ws.row_dimensions[ligne].height = hauteur
    X.name_cell(wb, nom, ws, debut)


def _titre(ws, cellule: str, texte: str) -> None:
    """Titre secondaire de la Synthèse (18)."""
    ws[cellule] = texte
    ws[cellule].font = TITRE_SECONDAIRE
    ws.row_dimensions[ws[cellule].row].height = 30


def build_synthese(wb) -> None:
    ws = _sheet(wb, "Synthèse", S.SAPIN)
    X.widths(ws, {get_column_letter(c): 15 for c in range(2, 14)})  # grille de 12 colonnes (B à M)
    ws["B2"] = "Synthèse stratégique"
    ws["B2"].font = TITRE_SYNTHESE
    ws.row_dimensions[2].height = 42
    X.name_cell(wb, "txt_sous_titre", ws, "B3", "", NOTE_12)
    # Le message clé : la situation en une phrase, sur toute la largeur
    _bloc(wb, ws, "message_cle", 4, 2, 13, MESSAGE, HAUT_GAUCHE, None, 54)

    # 1. Chiffres clés : 6 cartes de 2 colonnes
    _titre(ws, "B6", "Chiffres clés")
    for i in range(N_KPI):
        c1 = 2 + 2 * i
        police = CHIFFRE_ALERTE if i + 1 in KPI_ALERTE else GRAND_CHIFFRE
        _bloc(wb, ws, f"kpi_{i + 1}_valeur", 7, c1, c1 + 1, police, CENTRE, FOND_CARTE, 48)
        _bloc(wb, ws, f"kpi_{i + 1}_libelle", 8, c1, c1 + 1, LIBELLE, CENTRE, FOND_CARTE, 80)

    # 2. Ce qu'il faut retenir : 6 cartes de 4 colonnes, sur deux rangées
    _titre(ws, "B10", "Ce qu'il faut retenir")
    for i in range(N_INSIGHTS):
        ligne = 11 if i < 3 else 15
        c1 = 2 + 4 * (i % 3)
        bandeau = PatternFill(fill_type="solid", start_color=COULEURS_INSIGHTS[i], end_color=COULEURS_INSIGHTS[i])
        _bloc(wb, ws, f"retenir_{i + 1}_theme", ligne, c1, c1 + 3, THEME, GAUCHE, bandeau, 24)
        _bloc(wb, ws, f"retenir_{i + 1}_chiffre", ligne + 1, c1, c1 + 3, CHIFFRE_INSIGHT, GAUCHE, FOND_CARTE, 34)
        _bloc(wb, ws, f"retenir_{i + 1}_texte", ligne + 2, c1, c1 + 3, TEXTE, HAUT_GAUCHE, FOND_CARTE, 100)
        # un liseré blanc sépare deux cartes voisines
        for lg in range(ligne, ligne + 3):
            ws.cell(lg, c1 + 3).border = Border(right=Side(style="thick", color="FFFFFF"))

    # 3. Prévision du mois suivant : verdict par groupe (tableau, texte en 12)
    _titre(ws, "B19", "Prévision du mois suivant")
    head, nxt = X.add_table(ws, "T_Verdict", 20, 2, H_VERDICT, capacity=len(config.GROUP_CODES), wrap=True,
                            comment="Méthode choisie sur 12 mois, puis jugée sur les 12 mois suivants. Fourchettes à "
                                    "80 % laissées vides : aucune couverture démontrée. Détail : feuille Prévision.",
                            formats=TXT)
    ws.cell(20, 2).font = NOTE_12
    for col in range(2, 2 + len(H_VERDICT)):
        ws.cell(head, col).font = ENTETE_12
        ws.cell(head + 1, col).font = TEXTE  # ligne modèle : son style est recopié sur chaque ligne remplie
    ws.row_dimensions[head].height = 48
    ws.row_dimensions[head + 1].height = 34
    # Mise en évidence des groupes où un modèle fait mieux que la règle simple (verdict « Gain observé ... »)
    zone = f"F{head + 1}:F{head + len(config.GROUP_CODES)}"
    ws.conditional_formatting.add(zone, FormulaRule(formula=[f'LEFT($F{head + 1},4)="Gain"'],
                                                    font=Font(color=S.SAUGE, bold=True),
                                                    fill=PatternFill(fill_type="solid", start_color=S.BRUME,
                                                                     end_color=S.BRUME)))

    # 4. Recommandations : l'axe à gauche, l'action précise et chiffrée à droite
    _titre(ws, f"B{nxt}", "Recommandations")
    for i in range(N_RECOS):
        ligne = nxt + 1 + i
        _bloc(wb, ws, f"reco_{i + 1}_axe", ligne, 2, 4, TEXTE_GRAS, GAUCHE, FOND_CARTE, 66, LIGNE_BAS)
        _bloc(wb, ws, f"reco_{i + 1}_texte", ligne, 5, 13, TEXTE, GAUCHE, None, None, LIGNE_BAS)

    # 5. Limites : ce que les données ne permettent pas de conclure
    debut = nxt + N_RECOS + 2
    _titre(ws, f"B{debut}", "Limites")
    for texte, c1, c2 in [("Limite", 2, 4), ("Conséquence sur les conclusions", 5, 9), ("Comment la lever", 10, 13)]:
        ws.merge_cells(start_row=debut + 1, start_column=c1, end_row=debut + 1, end_column=c2)
        cell = ws.cell(debut + 1, c1, texte)
        cell.font, cell.fill, cell.alignment = ENTETE_12, S.HEAD_FILL, S.HEAD_ALIGN
        for col in range(c1, c2 + 1):
            ws.cell(debut + 1, col).fill = S.HEAD_FILL
    ws.row_dimensions[debut + 1].height = 26
    for i in range(N_LIMITES):
        ligne = debut + 2 + i
        _bloc(wb, ws, f"limite_{i + 1}_limite", ligne, 2, 4, TEXTE_GRAS, GAUCHE, FOND_CARTE, 50, LIGNE_BAS)
        _bloc(wb, ws, f"limite_{i + 1}_consequence", ligne, 5, 9, TEXTE, GAUCHE, None, None, LIGNE_BAS)
        _bloc(wb, ws, f"limite_{i + 1}_levee", ligne, 10, 13, TEXTE, GAUCHE, None, None, LIGNE_BAS)


def build_prevision(wb) -> None:
    """Feuille Prévision (contenu de Jeanne) : prévision, valeur métier, régularité, backtest, détail."""
    ws = _sheet(wb, "Prévision", S.C_CORAL)
    X.title(ws, "Peut-on prévoir le mois suivant ?")
    X.name_cell(wb, "txt_prevision_note", ws, "B3", "", S.NOTE)
    X.widths(ws, {"B": 12, "C": 30, "D": 30, "E": 16, "F": 16, "G": 16, "H": 16, "I": 16})
    head, nxt = X.add_table(ws, "T_Prevision", 5, 2, H_PREVISION, capacity=8, formats=TXT,
                            heading="Prévision du mois suivant",
                            comment="Méthode recommandée après validation. Les colonnes à 80 % restent vides : "
                                    "aucune couverture à 80 % n’a été démontrée.")
    head, nxt = X.add_table(ws, "T_Valeur", nxt, 2, H_VALEUR, capacity=8, formats=TXT,
                            heading="Valeur pour le métier : erreur évitée par rapport à la meilleure règle simple",
                            comment="Réduction d’erreur divisée par les ventes quotidiennes moyennes de validation. "
                                    "Équivalent en jours de ventes, pas économie de stock constatée. Une valeur "
                                    "négative indique une dégradation.")
    head, nxt = X.add_table(ws, "T_Regularite", nxt, 2, H_REGULARITE, capacity=8,
                            heading="Régularité de la demande et stock de sécurité",
                            comment="Coef. de variation = écart-type / moyenne : plus il est haut, plus il faut de "
                                    "stock de sécurité. P95 = niveau dépassé 1 jour sur 20.")
    head, nxt = X.add_table(ws, "T_Backtest", nxt, 2, H_BACKTEST, capacity=8,
                            heading="Résultats du backtest par groupe",
                            comment="MAE : erreur absolue moyenne en quantités par mois. MASE et p-valeurs non calculés : "
                                    "colonnes laissées vides. Les cinq méthodes sont détaillées dans le tableau ci-dessous.",
                            formats={**TXT, "Mois où le modèle fait mieux": S.FMT_INT,
                                     **{h: S.FMT_QTY for h in H_BACKTEST if h.startswith("MAE")}})
    X.add_table(ws, "T_Detail_Backtest", nxt, 2, H_DETAIL, heading="Détail : chaque prévision confrontée au mois réel",
                formats=TXT)


def build(wb) -> None:
    build_synthese(wb)
    build_prevision(wb)


def _remplir_cartes(wb, prefixe: str, df, champs: list[str]) -> None:
    """Écrit un DataFrame ligne par ligne dans les cartes nommées prefixe_<n>_<champ>."""
    for i, ligne in enumerate(df.itertuples(index=False), start=1):
        for champ, valeur in zip(champs, ligne):
            X.set_named(wb, f"{prefixe}_{i}_{champ}", valeur)


def fill(wb, c, t) -> dict[str, int]:
    X.set_named(wb, "txt_sous_titre", t["subtitle"])
    X.set_named(wb, "txt_prevision_note", t["fc_note"])
    if "kpi" in t:  # la synthèse croise les trois flux : remplie seulement quand tout est là
        _remplir_cartes(wb, "kpi", t["kpi"], ["valeur", "libelle"])
        _remplir_cartes(wb, "retenir", t["insights"], ["theme", "chiffre", "texte"])
        _remplir_cartes(wb, "reco", t["recommendations"], ["axe", "texte"])
        X.set_named(wb, "message_cle", t["message"])
    else:
        X.set_named(wb, "kpi_1_libelle", "En attente de l'intégration des trois flux")
    _remplir_cartes(wb, "limite", t["limits"], ["limite", "consequence", "levee"])
    lignes = {"T_Verdict": X.fill_table(wb, "T_Verdict", t["verdict"])}
    lignes.update(fill_prevision(wb, t))
    lignes.update(fill_detail_backtest(wb, t))
    lignes.update(fill_backtest(wb, t))
    return lignes


# ------------------------------------------------------------------ Prévision : remplissage (Jeanne)
def fill_prevision(wb, t):
    import pandas as pd

    verdicts = t["value"].set_index("Groupe")["Verdict : prévisible ?"]

    previsions = pd.DataFrame({
        "Groupe": t["fc_next"]["Groupe"],
        "Mois prévu": t["fc_next"]["Date"].dt.strftime("%Y-%m"),
        "Méthode retenue": t["fc_next"]["Méthode"],
        "Prévision": t["fc_next"]["Quantité prévue"],
        "Fourchette basse (80 %)": None,
        "Fourchette haute (80 %)": None,
        "Verdict": t["fc_next"]["Groupe"].map(verdicts)
    })

    return {
        "T_Prevision": X.fill_table(wb, "T_Prevision", previsions),
        "T_Valeur": X.fill_table(wb, "T_Valeur", t["value"]),
        "T_Regularite": X.fill_table(wb, "T_Regularite", t["regularity"])
    }

def fill_detail_backtest(wb, t):
    detail = t["fc_detail"].copy()

    detail["Date"] = detail["Date"].dt.strftime("%Y-%m")

    detail = detail.rename(columns={
        "Date": "Mois testé",
        "Prévu": "Prévision"
    })

    detail = detail[[
        "Groupe", "Mois testé", "Méthode",
        "Prévision", "Réel", "Erreur absolue"
    ]].sort_values(["Groupe", "Mois testé", "Méthode"])

    return {
        "T_Detail_Backtest": X.fill_table(
            wb, "T_Detail_Backtest", detail
        )
    }

def fill_backtest(wb, t):
    import pandas as pd
    from openpyxl.utils.cell import range_boundaries

    detail = t["fc_detail"]
    bilan = t["fc_summary"]
    valeur = t["value"].set_index("Groupe")

    scores = detail.pivot_table(
        index="Groupe",
        columns="Méthode",
        values="Erreur absolue",
        aggfunc="mean"
    )

    niveaux = detail.drop_duplicates(
        ["Date", "Groupe"]
    ).groupby("Groupe")["Réel"].mean()

    # Reprendre les intitulés exacts du modèle actuel
    feuille, tableau = X.tables(wb)["T_Backtest"]
    c1, ligne, c2, _ = range_boundaries(tableau.ref)
    colonnes = [
        feuille.cell(ligne, col).value
        for col in range(c1, c2 + 1)
    ]
    assert len(colonnes) == 19, "Structure du backtest inattendue."

    lignes = []

    for groupe, resultat in bilan.iterrows():
        methode = resultat["Méthode choisie avant validation"]
        reference = resultat["Référence choisie avant validation"]

        erreurs = detail.loc[
            detail["Groupe"] == groupe
        ].pivot(
            index="Date",
            columns="Méthode",
            values="Erreur absolue"
        )

        mois_gagnes = int(
            (erreurs[methode] < erreurs[reference]).sum()
        )

        lignes.append([
            groupe,
            niveaux[groupe],
            scores.loc[groupe, "Mois précédent"],
            scores.loc[groupe, "Même mois N-1"],
            scores.loc[groupe, "Moyenne 3 mois"],
            scores.loc[groupe, "Holt-Winters"],
            None, None, None, None,  # MASE non calculés
            reference,
            resultat["Erreur relative (%)"] / 100,
            resultat["Gain (%)"] / 100,
            mois_gagnes,
            None,  # p-valeur non calculée
            valeur.loc[groupe, "Verdict : prévisible ?"],
            valeur.loc[groupe, "Erreur évitée par le modèle (unités/mois)"],
            valeur.loc[groupe, "Équivalent en jours de vente"],
            valeur.loc[groupe, "Méthode recommandée"]
        ])

    return {
        "T_Backtest": X.fill_table(
            wb, "T_Backtest",
            pd.DataFrame(lignes, columns=colonnes)
        )
    }
